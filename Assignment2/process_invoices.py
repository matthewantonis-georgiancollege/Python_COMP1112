# Matthew Antonis
# Assignment 2
# Invoice Processor
# 200373088
# 07/18/2023

import os
from docx import Document
from openpyxl import Workbook

# Create a new Workbook
wb = Workbook()

# Grab the active worksheet from the workbook
ws = wb.active

# Add the headers to the Excel file
ws.append(["Invoice Number", "Total Quantity", "Subtotal", "Tax", "Total"])

# Specify the directory where the .docx files are located
directory = "Docs"

# Iterate over all files in the specified directory
for file in os.listdir(directory):
    # Combine the directory with the file name to create the full path
    full_path = os.path.join(directory, file)

    # Process only .docx files
    if file.endswith(".docx"):
        # Use the full path to open the .docx file
        doc = Document(full_path)

        # Initialize all variables
        invoice_number = ""
        total_quantity = 0
        subtotal = 0
        tax = 0
        total = 0
        # Flag to track if we are in the products section
        in_products_section = False

        # Iterate over all paragraphs in the document
        for paragraph in doc.paragraphs:
            # Check if the paragraph contains the invoice number
            if "INV" in paragraph.text:
                invoice_number = paragraph.text
            # Check if we have entered the products section
            elif "PRODUCTS" in paragraph.text:
                in_products_section = True
                # Skip the 'PRODUCTS' line and break down the paragraph into individual product lines
                product_lines = paragraph.text.split("\n")[1:]
                # Iterate over all product lines
                for product_line in product_lines:
                    # Make sure the line contains a colon (indicates product and quantity)
                    if ':' in product_line:
                        # Split the line into product name and quantity
                        product_name, product_quantity = product_line.split(':')
                        # Add the product quantity to the total quantity (remove trailing newline or whitespace characters before conversion)
                        total_quantity += int(product_quantity.strip())
            # Check if we have exited the products section and entered the financial details
            elif "SUBTOTAL" in paragraph.text:
                in_products_section = False
                # Break down the paragraph into individual financial lines
                financial_lines = paragraph.text.split('\n')
                # Iterate over all financial lines
                for line in financial_lines:
                    # Check for the various financial detail lines and extract their values
                    if "SUBTOTAL" in line:
                        subtotal = float(line.split(':')[1])
                    elif "TAX" in line:
                        tax = float(line.split(':')[1])
                    elif "TOTAL" in line:
                        total = float(line.split(':')[1])

        # Add the extracted details as a new row in the Excel file
        ws.append([invoice_number, total_quantity, subtotal, tax, total])

# Save the file
wb.save("./invoices.xlsx")
